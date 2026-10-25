"""Tailor + Verifier: build a resume for each queued job from the claim bank only.

1. Tailor (Claude): selects, orders and rewords claims for the job; every line cites claim IDs.
2. Verifier, code: cited IDs exist and belong where they're used; every master role appears
   (full or brief) so no gaps are created; dates, employers and contact details come from
   the claim bank, never from the model.
3. Verifier, Claude: each line is checked against the claims it cites; lines that add
   anything are removed.
4. Render: ATS-safe single-column .docx, plus .pdf through Word.

Run: .venv\\Scripts\\python -m jobhunter.tailor [--job ID] [--no-pdf]
"""
import argparse
import json
import re
import sys
from datetime import date

import yaml

from jobhunter import PRIVATE, claude, db, filters
from jobhunter.scout import load_settings

APPLICATIONS = PRIVATE / "applications"

LINE = {"type": "object", "required": ["text", "cites"], "properties": {
    "text": {"type": "string"}, "cites": {"type": "array", "items": {"type": "string"}, "minItems": 1}}}

PLAN_SCHEMA = {
    "type": "object",
    "required": ["summary", "skills", "roles", "recognition"],
    "properties": {
        "summary": {"type": "array", "items": LINE, "minItems": 2, "maxItems": 4},
        "skills": {"type": "array", "maxItems": 4, "items": {
            "type": "object", "required": ["label", "items"], "properties": {
                "label": {"type": "string"},
                "items": {"type": "array", "items": LINE, "maxItems": 9}}}},
        "roles": {"type": "array", "items": {
            "type": "object", "required": ["id", "mode", "note", "bullets"], "properties": {
                "id": {"type": "string"},
                "mode": {"enum": ["full", "brief"]},
                "note": {"type": "boolean"},
                "bullets": {"type": "array", "items": LINE}}}},
        "recognition": {"type": "array", "items": {"type": "string"}},
        "projects": {"type": "array", "items": LINE, "maxItems": 6},
    },
}

VERIFY_SCHEMA = {
    "type": "object", "required": ["lines"],
    "properties": {"lines": {"type": "array", "items": {
        "type": "object", "required": ["key", "ok", "problem"],
        "properties": {"key": {"type": "string"}, "ok": {"type": "boolean"}, "problem": {"type": "string"}}}}},
}

TAILOR_SYSTEM = """You tailor a resume for one job, using ONLY the candidate's claim bank. Return JSON only.

RULES
- Every line you write (summary sentence, skill item, bullet) cites the claim IDs it is based on.
- You may select, reorder, merge and reword claims, and use the job's vocabulary where it
  describes the same thing. You may NEVER add a fact: no new tools, techniques, numbers, results,
  scope, titles, regulations or responsibilities that the cited claims don't state.
- Keep each bullet as complete as its source: keep every tool, technique, system and document type
  the source lists (reorder them so the ones this job asks for come first). Never cut a bullet down
  to a fragment like "Documented deviations." A bullet is one full line or two; merging two related
  claims of the same role into one strong bullet is good.
- Bullets in a role cite only that role's claims (IDs that start with the role id). Summary lines
  cite sum.* and/or role claims. Skill items cite sk.* claims.
- Include EVERY role from the claim bank exactly once, by id. Use mode "full" for roles relevant to
  this job (bullets ordered by relevance to the job, at most {max_bullets} each) and mode "brief" for
  roles that aren't (no bullets: the resume then shows only title, employer and dates).
  The role "relocation" is always brief. Set note=true on a brief role only if a short
  "Experience outside the scope of this role" line would genuinely read better; usually false.
- Summary: 2-4 sentences leading with the strengths this job values most.
- Skills: 3-4 groups with job-relevant labels, most relevant items first.
- recognition: the rec.* IDs worth listing for this job (may be empty).
- projects: lines for the "AI & Technology Projects" section, citing only ai.* claims. Include them
  when they help this job; for lab jobs usually 0-2 (e.g. AI-assisted GxP documentation).
- Plain, professional US resume English. No first person. No contact details, dates or employer
  names in your text (those are filled in from the claim bank)."""

ALT_GUIDANCE = """

THIS IS AN ALTERNATIVE-TRACK JOB (outside the lab: AI implementation, project management, technical
writing, consulting, informatics, operations or QA testing).
- Lead the summary with the transferable strengths this job values: managing multiple labs and teams,
  priorities and decisions, assets, vendors and supply, GxP documentation, and hands-on AI work.
- Translate lab wording into general business language where the cited claim supports it (e.g.
  "instrument lifecycle" -> "equipment lifecycle"), but never add scope, numbers or results.
- Condense pure bench-technique detail; keep regulated-environment and documentation experience.
- Use 3-6 projects lines. Skills groups should feature AI & technology and management first.
- Roles that are relevant here may differ from lab jobs (e.g. the EA software QA tester role for QA
  testing jobs)."""

VERIFY_SYSTEM = """You check resume lines for faithfulness. For each line you get its text and the
source claims it cites. ok=true only if every fact in the line (tools, techniques, numbers, scope,
results, regulations, responsibilities, level of ownership) is stated or directly implied by the
cited sources. Rewording and the job's vocabulary for the same thing are fine. If anything is added
or exaggerated, ok=false and name it in problem (short). Return JSON only, one entry per key."""


def claims_index(bank):
    """{claim id: text} for every claim, plus {role id: role} for roles."""
    index = {}
    for s in bank["summary"]:
        index[s["id"]] = s["text"]
    for group in bank["skills"].values():
        for i in group:
            index[i["id"]] = i["text"]
    for r in bank["roles"]:
        for b in r["bullets"]:
            index[b["id"]] = b["text"]
    for e in bank["education"]:
        index[e["id"]] = e["degree"]
    for x in bank["recognition"] + bank["languages"] + bank.get("projects", {}).get("items", []):
        index[x["id"]] = x["text"]
    return index, {r["id"]: r for r in bank["roles"]}


def bank_text(bank):
    lines = ["SUMMARY CLAIMS"] + [f"[{s['id']}] {s['text']}" for s in bank["summary"]]
    lines.append("\nSKILL CLAIMS")
    for group, items in bank["skills"].items():
        lines.append(f"{group}: " + "; ".join(f"[{i['id']}] {i['text']}" for i in items))
    lines.append("\nROLES (id: title | employer | dates)")
    for r in bank["roles"]:
        lines.append(f"\nrole {r['id']}: {r['title']} | {r.get('employer') or ''} | {r['start']} to {r['end']}")
        lines += [f"  [{b['id']}] {b['text']}" for b in r["bullets"]]
    projects = bank.get("projects")
    if projects:
        lines.append(f"\nPROJECT CLAIMS ({projects['heading']}, {projects['start']} to {projects['end']})")
        lines += [f"[{p['id']}] {p['text']}" for p in projects["items"]]
    lines.append("\nRECOGNITION")
    lines += [f"[{x['id']}] {x['text']}" for x in bank["recognition"]]
    return "\n".join(lines)


def job_text(job, screen):
    sections = filters.extract_sections(job["description"], limit=3000)
    return "\n".join([
        f"JOB: {job['title']} at {job['employer']}",
        f"What the job is: {screen.get('summary', '')}",
        f"Screener's matches: {'; '.join(screen.get('matched', []))}",
        f"Screener's gaps: {'; '.join(screen.get('gaps', []))}",
        f"Duties:\n{sections['duties'] or ''}",
        f"Requirements:\n{sections['requirements'] or ''}",
    ])


TRUNCATION = 0.6  # a bullet shorter than this share of its sources is put back in full


def restore(bullet, index, problems):
    """Undo over-trimming: a bullet that drops most of its source falls back to the source text.

    The sources are Alex's own resume lines, so the full text is always safe to use.
    """
    sources = [index[c] for c in bullet["cites"]]
    if len(bullet["text"]) < TRUNCATION * sum(len(s) for s in sources):
        problems.append(f"restored full claim text for: {bullet['text'][:60]}")
        bullet["text"] = " ".join(s.rstrip(".") + "." for s in sources)
    return bullet


def code_check(plan, bank, max_bullets=6):
    """Repair what code can repair; return (plan, problems found)."""
    index, roles = claims_index(bank)
    problems = []

    def valid(line, prefixes):
        bad = [c for c in line["cites"] if c not in index or not c.startswith(prefixes)]
        if bad:
            problems.append(f"dropped line citing {bad}: {line['text'][:60]}")
        return not bad and line["text"].strip()

    plan["summary"] = [l for l in plan["summary"] if valid(l, ("sum.", "ai.") + tuple(f"{r}." for r in roles))]
    plan["projects"] = [restore(l, index, problems) for l in plan.get("projects", []) if valid(l, ("ai.",))]
    for group in plan["skills"]:
        group["items"] = [l for l in group["items"] if valid(l, ("sk.",))]
    plan["skills"] = [g for g in plan["skills"] if g["items"]]

    seen = {}
    for r in plan["roles"]:
        if r["id"] not in roles:
            problems.append(f"unknown role {r['id']}")
            continue
        seen[r["id"]] = r
        r["bullets"] = [restore(b, index, problems) for b in r["bullets"] if valid(b, (f"{r['id']}.",))]
        r["bullets"] = r["bullets"][:max_bullets]
        if r["id"] == "relocation":
            r.update(mode="brief", bullets=[], note=False)
        if r["mode"] == "brief":
            r["bullets"] = []
    for rid in roles:  # every master role appears, so the timeline keeps no new gaps
        if rid not in seen:
            problems.append(f"role {rid} missing; added as brief")
            seen[rid] = {"id": rid, "mode": "brief", "note": False, "bullets": []}
    # Roles always render newest first, whatever order the model chose.
    order = sorted(roles.values(), key=lambda r: str(r["start"]), reverse=True)
    plan["roles"] = [seen[r["id"]] for r in order]
    plan["recognition"] = [x for x in plan["recognition"] if x in index and x.startswith("rec.")]
    return plan, problems


def ai_check(conn, plan, bank, model):
    """Ask Claude to check each line against its sources; drop lines that add anything."""
    index, _ = claims_index(bank)
    lines = {}
    for i, l in enumerate(plan["summary"]):
        lines[f"summary.{i}"] = l
    for g, group in enumerate(plan["skills"]):
        for i, l in enumerate(group["items"]):
            lines[f"skills.{g}.{i}"] = l
    for r in plan["roles"]:
        for i, l in enumerate(r["bullets"]):
            lines[f"{r['id']}.{i}"] = l
    for i, l in enumerate(plan.get("projects", [])):
        lines[f"projects.{i}"] = l
    payload = "\n\n".join(
        f"key={k}\nLINE: {l['text']}\nSOURCES:\n" + "\n".join(f"- {index[c]}" for c in l["cites"])
        for k, l in lines.items())
    result = claude.run(conn, "verify", VERIFY_SYSTEM, payload, VERIFY_SCHEMA, model=model, jobs=1)
    rejected = {x["key"]: x["problem"] for x in result["lines"] if not x["ok"]}
    for key, line in lines.items():
        line["drop"] = key in rejected
    plan["summary"] = [l for l in plan["summary"] if not l.pop("drop")]
    for group in plan["skills"]:
        group["items"] = [l for l in group["items"] if not l.pop("drop")]
    for r in plan["roles"]:
        r["bullets"] = [l for l in r["bullets"] if not l.pop("drop")]
        if r["mode"] == "full" and not r["bullets"]:
            r["mode"] = "brief"
    plan["projects"] = [l for l in plan.get("projects", []) if not l.pop("drop")]
    return rejected


# --- Rendering -------------------------------------------------------------------------

def _month(value):
    if value == "present":
        return "Present"
    s = str(value)
    if len(s) == 4:
        return s
    y, m = s.split("-")
    return date(int(y), int(m), 1).strftime("%b %Y")


def render_docx(plan, bank, path):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Inches(0.75))
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Arial", Pt(10)
    normal.paragraph_format.space_after = Pt(2)

    def para(text="", size=10, bold=False, align=None, before=0, after=2):
        p = doc.add_paragraph()
        p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(before), Pt(after)
        if align:
            p.alignment = align
        if text:
            run = p.add_run(text)
            run.font.size, run.bold, run.font.name = Pt(size), bold, "Arial"
        return p

    def heading(text):
        para(text.upper(), size=11, bold=True, before=10, after=3)

    def bullet(text):
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(1)
        run = p.add_run(text)
        run.font.size, run.font.name = Pt(9.5), "Arial"

    h = bank["header"]
    para(h["name"], size=18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
    para(h["credentials"], size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
    para(" | ".join([h["phone"], h["email"], h["location"]]), size=10, align=WD_ALIGN_PARAGRAPH.CENTER, after=4)

    heading("Professional Summary")
    para(" ".join(l["text"] for l in plan["summary"]), size=10)

    heading("Core Competencies")
    for g in plan["skills"]:
        p = para(size=9.5, after=1)
        r = p.add_run(f"{g['label']}: ")
        r.bold, r.font.size = True, Pt(9.5)
        r = p.add_run(", ".join(l["text"] for l in g["items"]))
        r.font.size = Pt(9.5)

    heading("Professional Experience")
    roles = {r["id"]: r for r in bank["roles"]}
    for item in plan["roles"]:
        r = roles[item["id"]]
        dates = f"{_month(r['start'])} – {_month(r['end'])}"
        if item["id"] == "relocation":
            para(f"{r['title']}  |  {dates}", size=10, bold=True, before=6, after=0)
            para(r["bullets"][0]["text"], size=9.5)
            continue
        para(f"{r['title']}  |  {dates}", size=10, bold=True, before=6, after=0)
        where = "  |  ".join(x for x in (r.get("employer"), r.get("location"), r.get("mode")) if x)
        para(where, size=9.5, after=1)
        if item["mode"] == "full":
            for b in item["bullets"]:
                bullet(b["text"])
        elif item.get("note"):
            para("Experience outside the scope of this role.", size=9.5)

    projects = bank.get("projects")
    if projects and plan.get("projects"):
        heading(projects["heading"])
        para(f"{_month(projects['start'])} – {_month(projects['end'])}", size=9.5, after=1)
        for l in plan["projects"]:
            bullet(l["text"])

    heading("Education")
    for e in bank["education"]:
        years = f"{_month(e['start'])} – {_month(e['end'])}"
        para(f"{e['degree']}  |  {years}", size=10, bold=True, after=0)
        para(f"{e['school']}  |  {e['location']}  |  {e.get('note', '')}", size=9.5)

    recs = {x["id"]: x["text"] for x in bank["recognition"]}
    if plan["recognition"]:
        heading("Recognition & Publications")
        for rid in plan["recognition"]:
            bullet(recs[rid])

    heading("Languages")
    para("  |  ".join(l["text"] for l in bank["languages"]), size=9.5)
    doc.save(path)


def plain_text(plan, bank):
    """The resume as text for the dashboard preview."""
    roles = {r["id"]: r for r in bank["roles"]}
    out = ["SUMMARY", " ".join(l["text"] for l in plan["summary"]), "", "COMPETENCIES"]
    out += [f"{g['label']}: " + ", ".join(l["text"] for l in g["items"]) for g in plan["skills"]]
    out += ["", "EXPERIENCE"]
    for item in plan["roles"]:
        r = roles[item["id"]]
        out.append(f"{r['title']} | {r.get('employer') or ''} | {_month(r['start'])} – {_month(r['end'])}")
        out += [f"  • {b['text']}" for b in item["bullets"]]
    if plan.get("projects"):
        out += ["", bank["projects"]["heading"].upper()] + [f"  • {l['text']}" for l in plan["projects"]]
    return "\n".join(out)


def slug(text):
    return re.sub(r"[^A-Za-z0-9]+", "_", text).strip("_")[:40]


def prepare(conn, job, bank, config, make_pdf=True, rerender=False):
    """Tailor, verify and render one job. rerender=True reuses the stored plan (no Claude calls)."""
    t = config["tailor"]
    if rerender:
        plan = json.loads(job["resume"])
        plan, code_problems = code_check(plan, bank, t["max_bullets_per_role"])
        plan["verification"]["code"] += code_problems
    else:
        screen = json.loads(job["screen"])
        system = TAILOR_SYSTEM.replace("{max_bullets}", str(t["max_bullets_per_role"]))
        if job.get("track") == "alt":
            system += ALT_GUIDANCE
        prompt = f"CLAIM BANK\n{bank_text(bank)}\n\n{job_text(job, screen)}"
        plan = claude.run(conn, "tailor", system, prompt, PLAN_SCHEMA, model=t["model"], jobs=1)
        plan, code_problems = code_check(plan, bank, t["max_bullets_per_role"])
        rejected = ai_check(conn, plan, bank, t["verify_model"])
        plan["verification"] = {"code": code_problems, "rejected": rejected}

    folder = APPLICATIONS / f"{date.today().isoformat()}_{job['id']}_{slug(job['employer'])}"
    folder.mkdir(parents=True, exist_ok=True)
    name = f"Alex_Stoicu_Resume_{slug(job['employer'])}"
    docx_path = folder / f"{name}.docx"
    render_docx(plan, bank, docx_path)
    if make_pdf:
        from docx2pdf import convert
        convert(str(docx_path), str(folder / f"{name}.pdf"))
    plan["text"] = plain_text(plan, bank)
    plan["file"] = name

    conn.execute("UPDATE jobs SET resume = ?, resume_dir = ?, prepared_at = ?, status = 'prepared' WHERE id = ?",
                 (json.dumps(plan), str(folder), db.now(), job["id"]))
    v = plan["verification"]
    db.log_event(conn, job["id"], "prepared",
                 f"{len(v['code'])} code fixes, {len(v['rejected'])} lines removed by verification")
    conn.commit()
    return plan


def main(argv=None):
    parser = argparse.ArgumentParser()
    parser.add_argument("--job", type=int, help="prepare one job by id (any status)")
    parser.add_argument("--no-pdf", action="store_true")
    parser.add_argument("--rerender", action="store_true", help="re-apply code checks and re-render the stored plan")
    args = parser.parse_args(argv)

    config, _, _ = load_settings()
    with open(PRIVATE / "claims.yaml", encoding="utf-8") as f:
        bank = yaml.safe_load(f)
    conn = db.connect()
    if args.job:
        jobs = [dict(r) for r in conn.execute("SELECT * FROM jobs WHERE id = ?", (args.job,))]
    else:
        jobs = [dict(r) for r in conn.execute("SELECT * FROM jobs WHERE status = 'queued' ORDER BY rank_score DESC")]
    for job in jobs:
        try:
            plan = prepare(conn, job, bank, config, make_pdf=not args.no_pdf, rerender=args.rerender)
            v = plan["verification"]
            print(f"prepared {job['id']} {job['employer']} | {job['title']}: "
                  f"{len(v['code'])} code fixes, {len(v['rejected'])} lines removed")
        except Exception as e:
            print(f"failed {job['id']} {job['title']}: {e}", file=sys.stderr)
    if not jobs:
        print("nothing queued")


if __name__ == "__main__":
    main()
